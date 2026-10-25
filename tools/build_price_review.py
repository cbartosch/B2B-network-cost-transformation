#!/usr/bin/env python3
"""The researched prices as a workbook a steward decides in, and back.

    python tools/build_price_review.py
    python tools/build_price_review.py --to-register docs/price-review.xlsx

The first reads samples/outside_in/price_review.json - written by
tools/review_prices.py inside the api container - with the decisions in
samples/price_decisions.csv, and writes docs/price-review.xlsx. Every price
on the card is a row: undecided ones first, each prefilled with the decision
the register holds for it, if any.

The second reads a completed workbook back into the register. Only rows with
a decision are taken; a register line the workbook does not decide is left
as it was. Then:

    docker compose exec api sh -c 'cd /app && PYTHONPATH=/app \
        python tools/approve_prices.py --by "name" \
            --decisions /src/samples/price_decisions.csv'

A completed workbook from before this tool reads back too: its "Your
decision" and "Why" columns are taken for Decision and Reason.
"""
import argparse
import csv
import json
import sys
from decimal import Decimal
from pathlib import Path

REVIEW = Path("samples/outside_in/price_review.json")
REGISTER = Path("samples/price_decisions.csv")
SUMMARY = Path("samples/outside_in/summary.csv")
WORKBOOK = Path("docs/price-review.xlsx")
FIELDS = ["id", "decision", "base", "currency", "band_low", "band_high",
          "kind", "reason"]
DECISIONS = ("APPROVE", "REJECT", "HOLD")

#: Headers a completed workbook may use, mapped to the ones read back. The
#: first review was decided in a layout with its own names for both.
ALIASES = {"Your decision": "Decision", "Why": "Reason",
           "Kind of price": "Kind"}


def cents(value) -> str:
    return str(Decimal(str(value)).quantize(Decimal("0.01")))


def amount(value) -> str:
    """A typed band as the register writes it: 60, not 60.00."""
    d = Decimal(str(value))
    return str(int(d)) if d == d.to_integral_value() else cents(d)


def card_band(base) -> tuple:
    """The band approve_price gives a price nobody stated one for."""
    base = Decimal(str(base))
    return (base * Decimal("0.75"), base * Decimal("1.25"))


def register_line(row: dict, previous: dict | None) -> dict | None:
    """One register line from one decided workbook row, or None if undecided.

    A band left as the workbook's formula is the card's width, written
    blank; a band somebody typed is the band they decided.
    """
    decision = str(row.get("Decision") or "").strip().upper()
    if not decision:
        return None
    if decision not in DECISIONS:
        raise ValueError(f"{row.get('Price')}: decision must be one of "
                         f"{DECISIONS}, not {row.get('Decision')!r}")
    typed = row.get("band_typed") and row.get("Band low") is not None
    kind = row.get("Kind") or (previous or {}).get("kind") or ""
    return {"id": row["Price"], "decision": decision,
            "base": cents(row["Researched price"]),
            "currency": row["Currency"],
            "band_low": amount(row["Band low"]) if typed else "",
            "band_high": amount(row["Band high"]) if typed else "",
            "kind": str(kind).strip().upper().replace(" ", "_"),
            "reason": " ".join(str(row.get("Reason") or "").split())}


def rows_to_register(rows: list, register: dict) -> dict:
    """The register with every decided workbook row applied over it."""
    out = dict(register)
    for row in rows:
        line = register_line(row, register.get(row["Price"]))
        if line is not None:
            out[line["id"]] = line
    return out


def load_register(path=REGISTER) -> dict:
    if not Path(path).exists():
        return {}
    with open(path, encoding="utf-8") as f:
        return {line["id"]: line for line in csv.DictReader(f)}


def write_register(register: dict, path=REGISTER):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for row_id in sorted(register):
            w.writerow({k: register[row_id].get(k, "") for k in FIELDS})


def read_workbook(path) -> list:
    """The Prices sheet as dicts under the current header names."""
    from openpyxl import load_workbook

    values = load_workbook(path, data_only=True)["Prices"]
    formulas = load_workbook(path)["Prices"]
    header = [ALIASES.get(c.value, c.value) for c in values[1]]
    band = header.index("Band low") + 1
    rows = []
    for r in range(2, values.max_row + 1):
        row = {h: values.cell(row=r, column=i + 1).value
               for i, h in enumerate(header)}
        if not row.get("Price"):
            continue
        cell = formulas.cell(row=r, column=band).value
        row["band_typed"] = not (isinstance(cell, str)
                                 and cell.startswith("="))
        rows.append(row)
    return rows


# ------------------------------------------------------------------ workbook
def build(review: dict, register: dict, summary: dict, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fx = {s: Decimal(r["fx_rate_to_eur"]) for s, r in summary.items()}

    def eur(effect):
        return sum((Decimal(e["delta"]) * fx[s] for s, e in effect.items()
                    if "delta" in e and s in fx), Decimal(0))

    body, bold = Font(name="Arial"), Font(name="Arial", bold=True)
    head = (Font(name="Arial", bold=True, color="FFFFFF"),
            PatternFill("solid", fgColor="1F3A5F"))
    inputs = (Font(name="Arial", color="0000FF"),
              PatternFill("solid", fgColor="FFF2CC"))
    state_fill = {"UNDECIDED": "FCE4D6", "APPROVED": "E2EFDA",
                  "REVOKED": "F2F2F2"}
    wrap, top = Alignment(wrap_text=True, vertical="top"), \
        Alignment(vertical="top")
    money_eur = '#,##0;(#,##0);"-"'

    def header(ws, names, widths):
        for i, (name, width) in enumerate(zip(names, widths), start=1):
            c = ws.cell(row=1, column=i, value=name)
            c.font, c.fill, c.alignment = head[0], head[1], wrap
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = ws.cell(row=2, column=2)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    ws = wb.active
    ws.title = "Prices"
    cols = [("Price", 32), ("Country", 8), ("Product", 15), ("Tier (Mbps)", 9),
            ("State", 18), ("Decision", 11), ("Reason", 50), ("Band low", 10),
            ("Band high", 10), ("Researched price", 11), ("Currency", 8),
            ("Researched, USD/month", 12), ("Rate it replaces, USD/month", 12),
            ("Ratio", 7), ("Row it replaces", 13), ("Flags", 30),
            ("Vendor / plan", 40), ("Source", 32), ("Source wording", 50),
            ("Source URL", 40), ("As of", 10), ("Researched for", 16),
            ("Effect of its approval, EUR/yr", 14), ("Circuits re-priced", 10),
            ("Largest effects", 38),
            ("Approving it re-prices lower tiers", 36),
            ("Other companies found", 30)]
    header(ws, [c[0] for c in cols], [c[1] for c in cols])
    dv = DataValidation(type="list", formula1='"Approve,Reject,Hold"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    order = {"UNDECIDED": 0, "APPROVED": 1, "REVOKED": 2}
    prices = sorted(review["prices"], key=lambda p: (
        order[p["state"]], p["country"], p["product"], p["mbps"]))
    for r, p in enumerate(prices, start=2):
        line = register.get(p["id"]) or {}
        same = (line and line.get("base") == cents(p["base"])
                and line.get("currency") == p["currency"])
        ev = p["evidence"]
        src = (ev.get("sources") or [{}])[0]
        flags = list(p["flags"])
        if line and not same:
            flags.append(f"register decided {line['base']} {line['currency']}")
        # A stated band is typed; the card's width stays the formula, so
        # reading the workbook back leaves it the card's width.
        current = (Decimal(p["low"]), Decimal(p["high"]))
        stated = p["state"] == "APPROVED" and [cents(v) for v in current] \
            != [cents(v) for v in card_band(p["base"])]
        if line.get("band_low") and same:
            band = (float(line["band_low"]), float(line["band_high"]))
        elif stated:
            band = (float(current[0]), float(current[1]))
        else:
            band = (f"=J{r}*0.75", f"=J{r}*1.25")
        top3 = sorted(((s, Decimal(e["delta"]) * fx[s])
                       for s, e in p["effect"].items()
                       if "delta" in e and s in fx), key=lambda t: t[1])[:3]
        values = [
            p["id"], p["country"], p["product"], p["mbps"],
            p["state"].lower() + (f" by {p['decided_by']}"
                                  if p.get("decided_by") else ""),
            line.get("decision", "").title() if same else None,
            line.get("reason") if same else None, band[0], band[1],
            float(p["base"]), p["currency"],
            float(p["usd"]) if p["usd"] else None,
            float(p["replaces"]["usd"]) if p["replaces"] else None,
            f'=IF(OR(L{r}="",M{r}=""),"",L{r}/M{r})',
            p["replaces"]["basis"] if p["replaces"] else "none - new tier",
            "; ".join(flags), ev.get("label") or "",
            "; ".join(filter(None, [src.get("publisher"), "/".join(
                filter(None, [src.get("source_class"),
                              src.get("how_read")]))])) or ev["from"],
            src.get("excerpt") or ev.get("excerpt") or "",
            src.get("source_url") or "", ev.get("as_of") or p["as_of"] or "",
            p["researched_by"] or "not recorded", float(eur(p["effect"])),
            sum(e.get("circuits", 0) for e in p["effect"].values()),
            "; ".join(f"{s} {v:+,.0f}" for s, v in top3),
            "; ".join(p["inversions"]) or "no",
            "; ".join(f"{o['slug']} {o['value']} {o['currency']}"
                      for o in p["other_observations"])]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = body
            c.alignment = wrap if col in (7, 16, 17, 19, 25, 26, 27) else top
        for col in (6, 7, 8, 9):
            ws.cell(row=r, column=col).font = inputs[0]
            ws.cell(row=r, column=col).fill = inputs[1]
        ws.cell(row=r, column=5).fill = PatternFill(
            "solid", fgColor=state_fill[p["state"]])
        dv.add(ws.cell(row=r, column=6))
        for col in (8, 9, 10, 12, 13):
            ws.cell(row=r, column=col).number_format = "#,##0.00"
        ws.cell(row=r, column=14).number_format = "0.00x"
        ws.cell(row=r, column=23).number_format = money_eur
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(prices) + 1}"

    ss = wb.create_sheet("Sets")
    header(ss, ["Country", "Product", "Approved", "Undecided",
                "Effect of the approvals, EUR/yr",
                "Effect of approving the undecided, EUR/yr"],
           [8, 15, 30, 30, 16, 18])
    for r, s in enumerate(review["sets"], start=2):
        for col, v in enumerate([
                s["country"], s["product"],
                ", ".join(i.split("-")[-2] for i in s["approved"]) or "-",
                ", ".join(i.split("-")[-2] for i in s["undecided"]) or "-",
                float(eur(s["effect_of_approvals"])),
                float(eur(s["effect_of_approving_undecided"]))], start=1):
            c = ss.cell(row=r, column=col, value=v)
            c.font = body
        for col in (5, 6):
            ss.cell(row=r, column=col).number_format = money_eur
    total = len(review["sets"]) + 2
    ss.cell(row=total, column=1, value="Total").font = bold
    for col in (5, 6):
        L = get_column_letter(col)
        c = ss.cell(row=total, column=col, value=f"=SUM({L}2:{L}{total - 1})")
        c.font, c.number_format = bold, money_eur

    cs = wb.create_sheet("By company")
    header(cs, ["Company", "Baseline, EUR/yr", "Effect of the approvals, EUR/yr",
                "Share of baseline",
                "Effect of approving the undecided, EUR/yr"],
           [22, 16, 18, 10, 20])
    per = {"effect_of_approvals": {}, "effect_of_approving_undecided": {}}
    for s in review["sets"]:
        for kind in per:
            for slug, e in s[kind].items():
                if "delta" in e and slug in fx:
                    per[kind][slug] = (per[kind].get(slug, Decimal(0))
                                       + Decimal(e["delta"]) * fx[slug])
    slugs = sorted(set(per["effect_of_approvals"])
                   | set(per["effect_of_approving_undecided"]),
                   key=lambda s: per["effect_of_approvals"].get(s, 0))
    for r, slug in enumerate(slugs, start=2):
        for col, v in enumerate([
                slug, float(summary[slug]["baseline_base_eur"]),
                float(per["effect_of_approvals"].get(slug, 0)),
                f'=IF(B{r}=0,"",C{r}/B{r})',
                float(per["effect_of_approving_undecided"].get(slug, 0))],
                start=1):
            cs.cell(row=r, column=col, value=v).font = body
        for col in (2, 3, 5):
            cs.cell(row=r, column=col).number_format = money_eur
        cs.cell(row=r, column=4).number_format = '0.0%;-0.0%;"-"'

    rm = wb.create_sheet("Read me", 0)
    rm.column_dimensions["A"].width = 28
    rm.column_dimensions["B"].width = 100
    states = {}
    for p in prices:
        states[p["state"]] = states.get(p["state"], 0) + 1
    for r, (k, v) in enumerate([
            ("Researched prices", None),
            ("What this is", f"Every researched price on the rate card - "
             f"{', '.join(f'{n} {s.lower()}' for s, n in sorted(states.items()))} "
             "- beside the rate it would replace, with its evidence and what "
             "its approval does to the sample baselines."),
            ("Deciding", "Prices sheet: Decision is Approve, Reject or Hold, "
             "with a Reason. Band low and high are the range a price is "
             "approved at: the card's +/-25% unless typed over. A decision "
             "the register already holds is filled in; undecided prices "
             "come first."),
            ("Applying", "python tools/build_price_review.py --to-register "
             "<this file> writes samples/price_decisions.csv; "
             "tools/approve_prices.py --decisions applies it. A decision "
             "names its figure, and is refused if the card now holds "
             "another."),
            ("Effects", "EUR per year on the L0 baselines, priced as the "
             "estimate prices, against the card as it stands: an approved "
             "price shows the effect it is having, an undecided one the "
             "effect it would have. Effects of single prices are not "
             "additive; the Sets sheet's are."),
            ("Flags", "Read from the evidence, never decisions: wording "
             "that suggests wholesale, consumer, a 'from' floor or a "
             "promotion; one source; read from a snippet; untraceable; no "
             "exchange rate; re-prices lower tiers - a cheaper upper tier "
             "prices every circuit beneath it."),
            ("Checked", f"The pricing reproduced all {review['reproduced']} "
             "stored L0 baselines to the cent before any effect was "
             "computed.")], start=1):
        a = rm.cell(row=r, column=1, value=k)
        a.font = Font(name="Arial", bold=True, size=14 if r == 1 else 10)
        a.alignment = top
        if v:
            b = rm.cell(row=r, column=2, value=v)
            b.font, b.alignment = body, wrap
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return states


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--review", default=str(REVIEW))
    ap.add_argument("--register", default=str(REGISTER))
    ap.add_argument("--out", default=str(WORKBOOK))
    ap.add_argument("--to-register", default="",
                    help="read this completed workbook back into the "
                         "register instead of writing one")
    args = ap.parse_args()

    register = load_register(args.register)
    if args.to_register:
        merged = rows_to_register(read_workbook(args.to_register), register)
        changed = sorted(i for i in merged if merged[i] != register.get(i))
        write_register(merged, args.register)
        counts = {}
        for line in merged.values():
            counts[line["decision"]] = counts.get(line["decision"], 0) + 1
        print(f"{args.register}: {len(merged)} decision(s) "
              f"{counts}; {len(changed)} changed by the workbook"
              + (f": {', '.join(changed)}" if changed else ""))
        return 0

    review = json.loads(Path(args.review).read_text("utf-8"))
    with open(SUMMARY, encoding="utf-8") as f:
        summary = {r["slug"]: r for r in csv.DictReader(f)}
    states = build(review, register, summary, Path(args.out))
    print(f"{args.out}: {len(review['prices'])} researched price(s) {states}, "
          f"{len(review['sets'])} sets")
    return 0


if __name__ == "__main__":
    sys.exit(main())
